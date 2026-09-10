import pandas as pd
import numpy as np
import re
import os
import glob
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import warnings

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

# Definir columnas de salida
target_cols = [
    "Class", "Type", "Subtype", "Grantor", "Grantee", "Date", "Recorded", 
    "Inst.Date", "Acknowledged", "Filed", "Booktype", "Book", "Page", 
    "Inst.No.", "County", "State", "Notes", "Requirements", "Essences", 
    "Flags", "Restrictions", "Warnings", "Reviews", "Files"
]

def clean_specific_dates(date_val):
    if pd.isna(date_val) or date_val == "": return ""
    
    # Si ya es un objeto de fecha de Excel, formatear directamente
    if isinstance(date_val, datetime) or isinstance(date_val, pd.Timestamp): 
        return date_val.strftime('%m/%d/%Y')
        
    d_str = str(date_val).replace('\n', ' ').strip()
    
    # ---> REGLA NUEVA PARA FECHAS INCOMPLETAS (ej. 11/__/1961) <---
    if '_' in d_str and re.search(r'\d', d_str):
        # Si tiene números y guiones bajos, devolver el string tal cual (eskimpar el formateo)
        return d_str

    eff_match_post = re.search(r'(?i)(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})\s*(?:EFF|EFFECTIVE)', d_str)
    eff_match_pre = re.search(r'(?i)(?:EFF|EFFECTIVE|EFF\.|EFF:)\s*(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})', d_str)
    
    if eff_match_post:
        d_str = eff_match_post.group(1)
    elif eff_match_pre:
        d_str = eff_match_pre.group(1)
    else:
        if "DOD" in d_str.upper(): 
            d_str = re.sub(r'(?i)\s*DOD\s*', '', d_str).strip()
        date_pattern = re.search(r'(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})', d_str)
        if date_pattern:
            d_str = date_pattern.group(1)
            
    try:
        parsed_date = pd.to_datetime(d_str, errors='coerce')
        if pd.notna(parsed_date): 
            return parsed_date.strftime('%m/%d/%Y')
    except:
        pass
    return d_str

def clean_book(val):
    if pd.isna(val) or str(val).strip() == '': return ''
    val_str = str(val).split('\n')[0].strip()
    val_str = re.sub(r'(?i)^\s*(OR|0R)\s*', '', val_str)
    return val_str.strip()

def clean_page(val):
    if pd.isna(val) or str(val).strip() == '': return ''
    return str(val).split('\n')[0].strip()

def split_bp(bp):
    if pd.isna(bp) or str(bp).strip() == '': return '', ''
    bp_str = str(bp).split('\n')[0].strip()
    bp_str = re.sub(r'(?i)^\s*(OR|0R)\s*', '', bp_str).strip()
    
    if '/' in bp_str:
        parts = bp_str.split('/', 1)
    elif '-' in bp_str:
        parts = bp_str.split('-', 1)
    elif '_' in bp_str:
        parts = bp_str.split('_', 1)
    else:
        parts = bp_str.split()
        if len(parts) > 1:
            return parts[0].strip(), parts[-1].strip()
        return bp_str, ''
        
    return parts[0].strip(), parts[1].strip()

def read_excel_dynamic_header(filepath):
    df_raw = pd.read_excel(filepath, header=None)
    header_idx = -1
    for i in range(min(30, len(df_raw))):
        row_vals = [str(x).upper() for x in df_raw.iloc[i].dropna().values]
        row_clean = re.sub(r'[^A-Z]', '', "".join(row_vals))
        if (("GRANTOR" in row_clean and "GRANTEE" in row_clean) or 
            "INSTRUMENTTITLE" in row_clean or 
            "IMGINSTRUMENT" in row_clean or
            "INSTNAME" in row_clean or
            "PARTYOFTHEFIRSTPART" in row_clean):
            header_idx = i
            break
            
    if header_idx != -1:
        df = df_raw.iloc[header_idx+1:].reset_index(drop=True)
        cols = [str(c).strip() if pd.notna(c) else f"Unnamed_{j}" for j, c in enumerate(df_raw.iloc[header_idx].values)]
        df.columns = cols
        return df
    else:
        return pd.DataFrame()

def merge_split_rows(df):
    df = df.dropna(how='all').reset_index(drop=True)
    if len(df) == 0 or len(df.columns) == 0: return df
    col_0_name = df.columns[0]
    
    c0_clean = df[col_0_name].astype(str).str.replace(r'[^A-Z]', '', regex=True).str.upper()
    mask_headers = c0_clean.isin(['INSTRUMENTTITLE', 'IMGINSTRUMENT', 'INSTRUMENT', 'INSTNAME'])
    df = df[~mask_headers].reset_index(drop=True)
    
    key_col = None
    for c in df.columns:
        c_clean = re.sub(r'[^A-Z]', '', str(c).upper())
        if c_clean in ['IMAGENO', 'IMGINSTRUMENT', 'BOOKPAGE', 'BK', 'BKPG', 'VOLPG', 'VOL', 'INST', 'INSTR', 'DOCUMENT', 'FILINGDATE', 'FILEDDATE', 'FILEDATE']:
            key_col = c
            break
    if not key_col: return df 
    df[key_col] = df[key_col].replace(r'^\s*$', np.nan, regex=True)
    group_ids = df[key_col].notna().cumsum()
    def join_text(x):
        valid_strings = [str(val).strip() for val in x.dropna() if str(val).strip() != '']
        return ' '.join(valid_strings)
    agg_dict = {col: join_text for col in df.columns}
    agg_dict[key_col] = 'first' 
    df_merged = df.groupby(group_ids).agg(agg_dict).reset_index(drop=True)
    return df_merged

def extract_valid_sheets(filepath):
    try:
        xls = pd.ExcelFile(filepath)
    except Exception as e:
        print(f"Error al abrir {filepath}: {e}")
        return []
        
    sheets_data = []
    
    state_map = {
        'TEXAS': 'TX', 'NEW MEXICO': 'NM', 'OKLAHOMA': 'OK', 
        'LOUISIANA': 'LA', 'COLORADO': 'CO'
    }
    
    for sheet in xls.sheet_names:
        df_raw = pd.read_excel(xls, sheet_name=sheet, header=None)
        if df_raw.empty or len(df_raw.columns) == 0:
            continue
            
        header_idx = -1
        for i in range(min(30, len(df_raw))):
            row_vals = [str(x).upper() for x in df_raw.iloc[i].dropna().values]
            row_clean = re.sub(r'[^A-Z]', '', "".join(row_vals))
            
            if (("GRANTOR" in row_clean and "GRANTEE" in row_clean) or 
                "PARTYOFTHEFIRSTPART" in row_clean or
                "INSTRUMENTTITLE" in row_clean or 
                "IMGINSTRUMENT" in row_clean or
                "INSTNAME" in row_clean or
                ("INSTRUMENT" in row_clean and "BK" in row_clean)):
                header_idx = i
                break
                
        if header_idx != -1:
            ext_county, ext_state = "", ""
            preamble_cells = df_raw.iloc[:header_idx].fillna('').astype(str).values.flatten()
            
            for cell in preamble_cells:
                for line in str(cell).split('\n'):
                    match = re.search(r'(?:^|[,|-])\s*([A-Za-z\s]+?)\s+(?:CO\.|COUNTY)(?:[\s,]+(TX|TEXAS|NM|NEW\s*MEXICO|OK|OKLAHOMA|LA|LOUISIANA|CO|COLORADO|[A-Z]{2}))?\b', line, re.IGNORECASE)
                    
                    if match:
                        c_name = match.group(1).strip().title()
                        c_words = c_name.split()
                        if len(c_words) > 2:
                            c_name = " ".join(c_words[-2:])
                            
                        ext_county = c_name
                        if match.group(2):
                            raw_state = match.group(2).strip().upper()
                            raw_state = re.sub(r'\s+', ' ', raw_state)
                            ext_state = state_map.get(raw_state, raw_state)
                        break
                if ext_county: break
                
            df = df_raw.iloc[header_idx+1:].reset_index(drop=True)
            cols = [str(c).strip() if pd.notna(c) else f"Unnamed_{j}" for j, c in enumerate(df_raw.iloc[header_idx].values)]
            df.columns = cols
            
            sheets_data.append((df, ext_county, ext_state))
            
    return sheets_data

def load_and_merge_excels():
    all_files = glob.glob("*.xlsx") + glob.glob("*.xls") # Incluir los xls antiguos
    if "Normalized_Runsheet_CD.xlsx" in all_files: all_files.remove("Normalized_Runsheet_CD.xlsx")
    if "Normalized_Runsheet_NCD.xlsx" in all_files: all_files.remove("Normalized_Runsheet_NCD.xlsx")
    if "Normalized_Runsheet_Master.xlsx" in all_files: all_files.remove("Normalized_Runsheet_Master.xlsx")
    
    df_list = []
    
    for file in all_files:
        if file.startswith("~$"): continue
        try:
            valid_sheets = extract_valid_sheets(file)
            if not valid_sheets:
                print(f"Skipping {file} (No se detectaron tablas de runsheet)")
                continue
                
            for df, fallback_county, fallback_state in valid_sheets:
                df = merge_split_rows(df)
                temp_df = pd.DataFrame(columns=target_cols)
                temp_df['Original_Instrument'] = "" 
                
                col_instr, col_bk, col_pg, col_bp, col_inst_no = None, None, None, None, None
                col_idate, col_fdate, col_recorded = None, None, None
                col_grantor, col_grantee = None, None
                col_county, col_legal, col_tract, col_remarks, col_notes = None, None, None, None, None
                col_lands, col_int_conv, col_int_ret = None, None, None
                col_acres, col_min_res, col_deed_plot, col_doc_link = None, None, None, None

                for c in df.columns:
                    c_clean = re.sub(r'[^A-Z]', '', str(c).upper())
                    
                    if c_clean in ['INSTRUMENT', 'IMGINSTRUMENT', 'INSTRUMENTTITLE', 'INSTRUMENTTYPE', 'INSTR', 'INST', 'INSTNAME', 'DOCUMENT']: col_instr = c
                    elif c_clean in ['BK', 'BOOK', 'VOL']: col_bk = c
                    elif c_clean in ['PG', 'PAGE']: col_pg = c
                    elif c_clean in ['BOOKPAGE', 'BKPG', 'VOLPG', 'BOOKPAGE']: col_bp = c
                    elif c_clean in ['EXECEFFDATE', 'INSTDATE', 'INSTRUMENTDATE', 'DATEEXECUTED', 'DATE', 'DOD']: col_idate = c
                    elif c_clean in ['FILINGDATE', 'FILEDDATE', 'FILEDATE', 'FILED', 'DOF']: col_fdate = c
                    elif c_clean in ['DOR', 'DATERECORDED']: col_recorded = c
                    elif c_clean in ['PARTYOFTHEFIRSTPART', 'GRANTOR', 'GRANTER']: col_grantor = c
                    elif c_clean in ['PARTYOFTHESECONDPART', 'GRANTEE']: col_grantee = c
                    elif c_clean in ['COUNTYESTATE', 'COUNTY']: col_county = c
                    elif c_clean in ['LEGAL', 'LEGALDESCRIPTION']: col_legal = c
                    elif c_clean == 'TRACT': col_tract = c
                    elif c_clean in ['REMARKS', 'COMMENTS']: col_remarks = c
                    elif c_clean in ['NOTES', 'COMMENTSNOTES']: col_notes = c
                    elif c_clean in ['LANDSCOVERED', 'LANDS']: col_lands = c
                    elif c_clean == 'INTERESTCONVEYED': col_int_conv = c
                    elif c_clean == 'INTERESTRETAINED': col_int_ret = c
                    elif c_clean in ['IMAGENO', 'INSTNO']: col_inst_no = c 
                    elif 'ACRES' in c_clean: col_acres = c
                    elif 'MINERALROYALTY' in c_clean or 'RESERVATION' in c_clean: col_min_res = c
                    elif 'DEEDPLOT' in c_clean: col_deed_plot = c
                    elif 'DOCUMENTLINK' in c_clean or 'LINK' in c_clean: col_doc_link = c

                if col_instr: temp_df['Original_Instrument'] = df[col_instr]
                if col_grantor: temp_df['Grantor'] = df[col_grantor]
                if col_grantee: temp_df['Grantee'] = df[col_grantee]
                if col_idate: temp_df['Inst.Date'] = df[col_idate].apply(clean_specific_dates)
                if col_fdate: temp_df['Filed'] = df[col_fdate].apply(clean_specific_dates)
                if col_recorded: temp_df['Recorded'] = df[col_recorded].apply(clean_specific_dates)
                if col_inst_no: temp_df['Inst.No.'] = df[col_inst_no]

                if col_bk and col_pg:
                    temp_df['Book'] = df[col_bk].apply(clean_book)
                    temp_df['Page'] = df[col_pg].apply(clean_page)
                elif col_bp:
                    temp_df['Book'] = df[col_bp].apply(lambda x: split_bp(x)[0])
                    temp_df['Page'] = df[col_bp].apply(lambda x: split_bp(x)[1])

                if col_county:
                    temp_df['County'] = df[col_county].apply(lambda x: str(x).split(',')[0].strip() if pd.notna(x) else '')
                    temp_df['State'] = df[col_county].apply(lambda x: str(x).split(',')[1].strip() if pd.notna(x) and ',' in str(x) else '')
                else:
                    temp_df['County'] = fallback_county
                    temp_df['State'] = fallback_state
                
                legal = df[col_legal] if col_legal else pd.Series(['']*len(df))
                tract = df[col_tract] if col_tract else pd.Series(['']*len(df))
                lands = df[col_lands] if col_lands else pd.Series(['']*len(df))
                int_conv = df[col_int_conv] if col_int_conv else pd.Series(['']*len(df))
                int_ret = df[col_int_ret] if col_int_ret else pd.Series(['']*len(df))
                remarks = df[col_remarks] if col_remarks else (df[col_notes] if col_notes else pd.Series(['']*len(df)))
                
                for idx in range(len(temp_df)):
                    leg_val = str(legal.iloc[idx]) if pd.notna(legal.iloc[idx]) and str(legal.iloc[idx]) != '-' else ''
                    trac_val = str(tract.iloc[idx]) if pd.notna(tract.iloc[idx]) and str(tract.iloc[idx]) != '-' else ''
                    land_val = str(lands.iloc[idx]) if pd.notna(lands.iloc[idx]) and str(lands.iloc[idx]) != '-' else ''
                    iconv_val = str(int_conv.iloc[idx]) if pd.notna(int_conv.iloc[idx]) and str(int_conv.iloc[idx]) != '-' else ''
                    iret_val = str(int_ret.iloc[idx]) if pd.notna(int_ret.iloc[idx]) and str(int_ret.iloc[idx]) != '-' else ''
                    instr_val = str(temp_df['Original_Instrument'].iloc[idx]) if pd.notna(temp_df['Original_Instrument'].iloc[idx]) else ''
                    
                    notes_parts = []
                    if leg_val: notes_parts.append(f"Legal: {leg_val}")
                    if trac_val: notes_parts.append(f"Tract: {trac_val}")
                    if land_val: notes_parts.append(f"Lands Covered: {land_val}")
                    if iconv_val: notes_parts.append(f"Interest Conveyed: {iconv_val}")
                    if iret_val: notes_parts.append(f"Interest Retained: {iret_val}")
                    
                    if col_acres and pd.notna(df[col_acres].iloc[idx]) and str(df[col_acres].iloc[idx]).strip():
                        notes_parts.append(f"Acres: {df[col_acres].iloc[idx]}")
                    if col_min_res and pd.notna(df[col_min_res].iloc[idx]) and str(df[col_min_res].iloc[idx]).strip():
                        notes_parts.append(f"Mineral/Royalty Reservation: {df[col_min_res].iloc[idx]}")
                    if col_deed_plot and pd.notna(df[col_deed_plot].iloc[idx]) and str(df[col_deed_plot].iloc[idx]).strip():
                        notes_parts.append(f"Deed Plot: {df[col_deed_plot].iloc[idx]}")
                    if col_doc_link and pd.notna(df[col_doc_link].iloc[idx]) and str(df[col_doc_link].iloc[idx]).strip():
                        notes_parts.append(f"Document Link: {df[col_doc_link].iloc[idx]}")
                    
                    if instr_val: 
                        notes_parts.append(f"Original Instrument: {instr_val}")
                        instr_upper = instr_val.upper()
                        
                        if len(instr_val) >= 75 or "CONSOLIDATED DT" in instr_upper or "SUPPLEMENT TO DT" in instr_upper:
                            notes_parts.append("Check instrument couldnt determine what it is form the runsheet import")
                        
                        if any(j in instr_upper for j in ["JUDG", "JUDGMENT"]):
                            notes_parts.append("Note: This document is a Judgment.")
                        if any(a in instr_upper for a in ["AGREE", "TRUST AGREE", "STIP", "JOA", "CONTRACT"]):
                            notes_parts.append("Note: This document is an Agreement.")
                    
                    temp_df.at[idx, 'Notes'] = "\n\n".join(notes_parts)
                    temp_df.at[idx, 'Essences'] = str(remarks.iloc[idx]) if pd.notna(remarks.iloc[idx]) and str(remarks.iloc[idx]) != '-' else ''

                df_list.append(temp_df)
            print(f"Archivo procesado exitosamente: {file}")
            
        except Exception as e:
            print(f"Error procesando el archivo {file}: {e}")
            
    if not df_list:
        df_empty = pd.DataFrame(columns=target_cols)
        df_empty['Original_Instrument'] = ""
        return df_empty
        
    final_df = pd.concat(df_list, ignore_index=True)
    
    final_df = final_df[
        ~(
            final_df['Grantor'].astype(str).str.strip().str.upper().eq('GRANTOR') |
            final_df['Original_Instrument'].astype(str).str.replace(r'[^A-Z]', '', regex=True).isin(['INSTRUMENTTITLE', 'INSTRUMENT', 'IMGINSTRUMENT', 'INSTNAME'])
        )
    ].reset_index(drop=True)
    
    valid_county = final_df['County'].replace('', np.nan).dropna()
    valid_state = final_df['State'].replace('', np.nan).dropna()
    global_county = valid_county.iloc[0] if not valid_county.empty else ""
    global_state = valid_state.iloc[0] if not valid_state.empty else ""
    final_df['County'] = global_county
    final_df['State'] = global_state
    
    return final_df

df_final = load_and_merge_excels()

def get_year(date_val):
    if pd.isna(date_val) or date_val == "": return None
    match = re.search(r'\d{4}', str(date_val))
    return int(match.group()) if match else None

def assign_booktype(row):
    year = get_year(row['Inst.Date'])
    if not year: year = get_year(row['Filed'])
    if year: return "DR" if year < 1990 else "OPR"
    return ""

if not df_final.empty:
    df_final['Booktype'] = df_final.apply(assign_booktype, axis=1)

# ---> DICCIONARIO ACTUALIZADO CON ABREVIATURAS RARAS DE LOS 60s/70s <---
SUBSTRING_RULES = {
    "ASSIGNMENT OF OIL AND GAS LEASE": "AOGL",
    "PARTIAL ASSIGNMENT OF OIL AND GAS LEASE": "PART AOGL",
    "ASSIGNMENT OF OIL, GAS, AND MINERAL LEASE": "AOGML",
    "PARTIAL ASSIGNMENT OF OIL, GAS, AND MINERAL LEASE": "PART AOGML",
    "ASSIGNMENT OF OVERRIDING ROYALTY INTEREST": "AORRI",
    "ASSIGNMENT OF OVERRIDING ROAYLTY INTEREST": "AORRI",
    "ASSIGNMENT OF DEED OF TRUST": "ASSIGN DT",
    "ASSIGNMENT OF PRODUCTION": "DT", 
    "ASSIGNMENT OF AS-EXTRACTED COLLATERAL": "DT", 
    "BILL OF SALE AND ASSIGNMENT": "ABOS",
    "ASSIGNMENT AND BILL OF SALE": "ABOS",
    "PARTIAL ASSIGNMENT": "PART ASSIGN",
    "ASSIGNMENT": "ASSIGN",
    "BLANKET MINERAL AND ROYALTY DEED": "MD",
    "MINERAL AND ROYALTY DEED": "MRD",
    "MINERAL/ROYALTY DEED": "MRD",
    "MINERAL AND ROAYLTY INTEREST DEED": "MRD",
    "CORRECTION MINERAL AND ROYALTY DEED": "MRD",
    "MINERAL DEED": "MD",
    "INDEPENDENT CO-EXECUTOR'S DISTRIBUTION DEED": "DIST DEED",
    "DISTRIBUTION DEED": "DIST DEED",
    "TRUSTEE'S DISTRIBUTION DEED": "DIST DEED",
    "DEED OF DISTRIBUTION": "DIST DEED",
    "AFFIDAVIT OF HEIRSHIP": "AFFT HEIR",
    "AFF OF HEIRSHIP": "AFFT HEIR",
    "AFFIDAVIT OF SUCCESSOR TRUSTEE": "AFFT SUCCESSOR TRST",
    "AFFIDAVIT": "AFFT",
    "SURFACE USE AND EASEMENT AGREEMENT": "AGREE",
    "EASEMENT": "EAS",
    "DEED OF TRUST": "DT",
    "QUIT CLAIM": "QCD",
    "GENERAL WARRANTY DEED": "GWD",
    "SPECIAL WARRANTY DEED": "SWD",
    "WARRANTY DEED WITH VENDOR'S LIEN": "WDVL",
    "WARRANTY DEED": "WD",
    "GIFT DEED": "GIFT DEED",
    "ADMINISTRATOR'S DEED": "ADMIN DEED",
    "EXECUTOR'S DEED": "EXEC DEED",
    "PARTITION DEED": "PART DEED",
    "PERSONAL REPRESENTATIVE'S DEED": "PR DEED",
    "RIGHT OF WAY DEED": "ROW DEED",
    "SHERIFF'S DEED": "SHRF DEED",
    "SUCCESSOR TRUSTEE'S DEED": "SUCCESSOR TRST DEED",
    "TRUSTEE'S DEED": "TRST DEED",
    "ROYALTY DEED": "RD",
    "DEED": "DEED",
    "BILL OF SALE": "BOS",
    "PARTIAL RELEASE OF LIEN": "RELL OF LIEN", 
    "RELEASE OF LIEN": "RELL OF LIEN",         
    "LIEN": "RELL OF LIEN",                    
    "RELEASE": "REL",
    "AMENDED AND RESTATED": "AMEND",
    "AMEND": "AMEND",
    "MEMORANDUM OF OIL AND GAS LEASE": "MEMO OGL",
    "MEMORANDUM OF OIL, GAS, AND MINERAL LEASE": "MEMO OGML",
    "OIL AND GAS LEASE": "OGL",
    "OIL & GAS LEASE": "OGL",                  
    "OIL, GAS, AND MINERAL LEASE": "OGML",
    "CC PROBATE": "PROBATE",                           
    "CC WILL": "PROBATE",                              
    "CC WILL / ORDER": "PROBATE",                      
    "WILL & PROBATE": "PROBATE",                       
    "CC/PROBATE": "PROBATE",                           
    "PROBATE": "PROBATE",
    "ESTATE": "PROBATE",
    "AGREEMENT": "AGREE",
    "CONVEYANCE": "CONV"
}

SORTED_SUBSTRING_RULES = sorted(SUBSTRING_RULES.items(), key=lambda x: len(x[0]), reverse=True)

# ---> ABREVIATURAS VIEJAS AÑADIDAS AQUÍ <---
EXACT_RULES = {
    r'\bAOGL\b': 'AOGL', r'\bAOGML\b': 'AOGML', r'\bAORRI\b': 'AORRI', r'\bASORI\b': 'AORRI',
    r'\bASSIGN DT\b': 'ASSIGN DT', r'\bABOS\b': 'ABOS', r'\bASSN\b': 'ASSIGN', r'\bASGN\b': 'ASSIGN',
    r'\bASN\b': 'ASSIGN', r'^AS$': 'ASSIGN', r'\bPART AOGL\b': 'PART AOGL', r'\bPART AOGML\b': 'PART AOGML',
    r'\bPART ASSIGN\b': 'PART ASSIGN', r'\bPART ASGN\b': 'PART ASSIGN', r'\bMD\b': 'MD', r'\bMRD\b': 'MRD', 
    r'\bDIST DEED\b': 'DIST DEED', r'\bAFFT HEIR\b': 'AFFT HEIR', r'\bAFF H\b': 'AFFT HEIR', r'\bAFFT\b': 'AFFT', 
    r'\bAGREE\b': 'AGREE', r'\bAGMT\b': 'AGREE', r'\bDOTO\b': 'DT', r'\bD OF TR\b': 'DT', r'\bDT\b': 'DT',
    r'\bQCD\b': 'QCD', r'\bWD\b': 'WD', r'\bCOR WD\b': 'WD', r'\bCORRECTION WD\b': 'WD', r'\bWD RECORD\b': 'WD', r'\bWD/ VL\b': 'WD',
    r'\bGWD\b': 'GWD', r'\bSWD\b': 'SWD', r'\bWDVL\b': 'WDVL', r'\bBOS\b': 'BOS', r'\bCONV\b': 'CONV', 
    r'\bREL\b': 'REL', r'\bREL LN\b': 'RELL OF LIEN', r'\bREL. OL\b': 'REL',
    r'\bAMD\b': 'AMEND', r'\bOGL\b': 'OGL', r'\bOGML\b': 'OGML', r'\bPROB\b': 'PROBATE', r'\bCC PROB\b': 'PROBATE',
    r'\bWILL\b': 'PROBATE', r'\bPATENT\b': 'PATENT', r'\bHEIRSHIP\b': 'AFFT HEIR', r'\bO&GL\b': 'OGL', r'\bOG&ML\b': 'OGML', r'\bMOGL\b': 'OGL',
    r'\bJDG\b': 'JUDGMENT'
}

def assign_subtype_and_type(instrument):
    if pd.isna(instrument): return "MISC", "MISC.", "NCD"
    
    instr = str(instrument).upper().strip() 
    subtype = ""
    
    for phrase, sub_val in SORTED_SUBSTRING_RULES:
        if phrase in instr:
            subtype = sub_val
            break
            
    if not subtype:
        for pattern, sub_val in EXACT_RULES.items():
            if re.search(pattern, instr):
                subtype = sub_val
                break
                
    if not subtype:
        subtype = instr.split('\n')[0].strip()

    conveyances = ["MD", "WD", "GWD", "SWD", "QCD", "DEED", "CONV", "BOS", "PROBATE", "AFFT HEIR", "MRD", "DIST DEED", "GIFT DEED", "ADMIN DEED", "EXEC DEED", "PART DEED", "PR DEED", "SHRF DEED", "SUCCESSOR TRST DEED", "TRST DEED", "RD", "ROW DEED", "WDVL", "PATENT"]
    assignments = ["ASSIGN", "ABOS", "AOGL", "AOGML", "PART ASSIGN", "PART AOGL", "PART AOGML", "AORRI", "ASSIGN DT", "BLM ASSIGN", "BLM TRANSFER", "TOD"]
    leases = ["OGL", "OGML"]
    releases = ["REL"]
    
    if any(c == subtype or f" {c} " in f" {subtype} " for c in conveyances) or "PATENT" in instr: return subtype, "Conveyance", "CD"
    if any(a == subtype or f" {a} " in f" {subtype} " for a in assignments): return subtype, "Assignment", "CD"
    if any(l == subtype or f" {l} " in f" {subtype} " for l in leases): return subtype, "Lease", "CD"
    if any(r == subtype or f" {r} " in f" {subtype} " for r in releases): return subtype, "Release", "CD"

    doc_class = "NCD"
    easements = ["EAS", "ROW", "ROW/EAS", "EASEMENT"]
    d_of_trust = ["DT", "MTG", "SUB", "DEED OF TRUST", "D OF TR"]
    liens = ["MML", "FTL", "STL", "HL", "AJ", "RELL OF LIEN", "LIEN"]
    agreements = ["AGREE", "TRUST AGREE", "STIP", "JOA", "RDO", "AGREEMENT", "CONTRACT"]
    judgements = ["JUDG", "JUDGMENT"]

    if any(d == subtype or f" {d} " in f" {subtype} " for d in d_of_trust) or "DEED OF TRUST" in instr or "D OF TR" in instr or r'\bDT\b' in instr or re.search(r'\bDT\b', instr):
        return "DT", "DEED OF TRUST", doc_class
    elif any(e == subtype or f" {e} " in f" {subtype} " for e in easements) or "EASEMENT" in instr:
        return "Easement", "Easement", doc_class
    elif any(j == subtype or f" {j} " in f" {subtype} " for j in judgements) or "JUDG" in instr:
        return "JUD", "MISC.", doc_class
    elif "AMEND" in subtype or "AMEND" in instr:
        return "AMENDM", "MISC.", doc_class
    elif any(a == subtype or f" {a} " in f" {subtype} " for a in agreements) or "AGREEMENT" in instr or "CONTRACT" in instr:
        return "AGREMNT", "AGRE", doc_class
    elif any(l == subtype or f" {l} " in f" {subtype} " for l in liens) or "LIEN" in instr:
        return "RELL OF LIEN", "Lien", doc_class
    else:
        return "MISC", "MISC.", doc_class

if not df_final.empty:
    mapped_values = df_final['Original_Instrument'].apply(assign_subtype_and_type)
    df_final['Subtype'] = mapped_values.apply(lambda x: x[0])
    df_final['Type'] = mapped_values.apply(lambda x: x[1])
    df_final['Class'] = mapped_values.apply(lambda x: x[2])
else:
    df_final['Subtype'] = ""
    df_final['Type'] = "MISC."
    df_final['Class'] = "NCD"

df_final = df_final[target_cols]

df_cd = df_final[df_final['Class'] == 'CD'].copy()
df_ncd = df_final[df_final['Class'] == 'NCD'].copy()

def save_formatted_excel(df, out_file):
    if df.empty:
        print(f"⚠️ No hay registros para generar {out_file}.")
        return
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Normalized Runsheet"

    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'), 
                    top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF'))

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            else:
                cell.border = border
                cell.alignment = Alignment(vertical='top', wrap_text=True)

    ws.column_dimensions['A'].width = 10 
    ws.column_dimensions['B'].width = 15 
    ws.column_dimensions['C'].width = 20 
    ws.column_dimensions['D'].width = 25 
    ws.column_dimensions['E'].width = 25 
    ws.column_dimensions['H'].width = 15 
    ws.column_dimensions['J'].width = 15 
    ws.column_dimensions['N'].width = 15 
    ws.column_dimensions['O'].width = 15 
    ws.column_dimensions['Q'].width = 50 
    ws.column_dimensions['S'].width = 50 

    wb.save(out_file)
    print(f"✅ Proceso terminado. Documento generado: {out_file} ({len(df)} registros)")

print("\n--- GENERANDO ARCHIVOS SEPARADOS ---")
save_formatted_excel(df_cd, "Normalized_Runsheet_CD.xlsx")
save_formatted_excel(df_ncd, "Normalized_Runsheet_NCD.xlsx")