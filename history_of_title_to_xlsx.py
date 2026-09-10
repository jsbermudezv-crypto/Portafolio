#!/usr/bin/env python3
"""
history_of_title_to_xlsx.py

Converts a "History of Title" style document (a narrative title-opinion exhibit, NOT a
tabular runsheet) into normalized spreadsheets ready to import into the land software.

USAGE
    python3 history_of_title_to_xlsx.py input.(txt|pdf) output_prefix [county] [state]

    The input may be the PDF itself or a plain-text export of it (auto-detected by
    extension). Feeding it the PDF is preferred: the script then reconstructs each line's
    original indentation from the word positions, which is what lets it tell a real item
    from a nested numbered sub-list (see EXPECTED INPUT SHAPE below).

    county/state are auto-detected from the document text; pass them on the command line
    only to override that detection, or to supply a value when nothing is found.

OUTPUT FILES
    <output_prefix>_HistoryOfTitle.xlsx   one row per item, 17 descriptive columns
    <output_prefix>_CD.xlsx              Conveyance / Assignment / Lease / Release
    <output_prefix>_NCD.xlsx             everything else (DT, Easement, Agreement, Lien...)
    <output_prefix>_MissingFields.txt    every row with a blank or uncertain field
    <output_prefix>_extracted.txt        (PDF input only) the text as the parser read it

EXPECTED INPUT SHAPE
    The .txt file should look like the plain-text body of the document, e.g.:

        A. Tract 1:

        1.  By Warranty Deed dated April 12, 1907, recorded in Volume 7, Page 9,
            Deed Records, H. A. Snively, et ux., conveyed Section 23 to J. L. Allred.

        2.  By Warranty Deed dated November 5, 1907, ...

        B. Tract 2:

        1.  ...

    - A TRACT header is a line starting "A.", "B.", "C." ... followed by "Tract ...:"
      at the LEFT MARGIN (column 0, or with the same small indent as item numbers).
    - An ITEM is a line starting "N." (a number + period) at that same left margin.
      Its text may wrap onto following lines -- those continuation lines are joined
      back into one paragraph automatically.
    - A FOOTNOTE is a line that starts with a bare number (no period) followed by its
      text, e.g. "12 We have assumed ..." -- these are pulled out separately and
      matched back to whichever item ends in that footnote number (e.g. "...RD.12").
    - Nested numbered sub-lists inside an item (e.g. a list of heirs "1. Jane Doe...
      2. John Doe...") are the one thing plain text can't disambiguate from real
      items the way the original PDF parser could (it used each line's physical
      left-edge position on the page). This script falls back to INDENTATION: a
      digit-led line is only treated as a real item/tract header if it starts at
      (or very near) the file's dominant left margin; anything indented further is
      treated as a continuation of the current item's text. If your .txt file has
      no preserved indentation at all, nested sub-lists will be mis-detected as
      items -- re-indent them by hand (a few spaces) before running this script.

THE "HISTORY OF TITLE" SHEET
    One row per item, with columns:
    Tract, Tract Description, Chain, Item #, Instrument Type, Date Executed,
    Effective Date, Volume, Page, Records Type, Deed Nickname, Grantor,
    Interest / Conveyance, Grantee, Footnote #, Footnote Text, Full Item Text.

    "Chain" increments every time an item's number is LOWER than the previous
    item's number within the same tract -- some title opinions restart numbering
    mid-tract for a second, independent chain of title (e.g. a separate mineral vs.
    royalty chain). Items that don't match the standard sentence pattern (probate
    narratives handled separately; anything else -- judgments, stipulations, etc.)
    are left with blank structured fields but their full paragraph is always kept
    in "Full Item Text", so no information is ever silently dropped.

    A Grantor or Grantee that cannot be determined is never left blank: it is written
    as "Unclear, see instrument" (or, for a multi-generation heirship narrative, as
    "Estate of X, Dec'd" / "Heirs of X, see instrument") and listed in the
    _MissingFields.txt report for review.
"""

import os
import sys
import re
from collections import Counter
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# =======================================================================================
# STAGE 1 -- split the raw text into Tracts -> Items (+ a separate footnote lookup)
# =======================================================================================

TRACT_PAT = re.compile(r'^([A-Z])\.\s+(Tract[^:]*):\s*$')
ITEM_PAT = re.compile(r'^(\d{1,3})\.\s+(.*)$')
# a footnote line: bare number + space + text, WITHOUT the period that a real item has
FOOTNOTE_PAT = re.compile(r'^(\d{1,3})\s+([A-Z(].*)$')

# lines to discard outright (repeating headers/footers/page numbers) -- extend this list
# to match whatever boilerplate your source document repeats on every page
NOISE_PATS = [
    re.compile(r'^\d{1,4}$'),                       # a bare page number on its own line
    re.compile(r'^Page \d+ of \d+$', re.I),
    re.compile(r'^Exhibit ["\']?[A-Z]["\']?.*$', re.I),
    re.compile(r'^[A-Za-z][A-Za-z .,&\'-]{2,50}:\s*\d{1,8}$'),  # Bates stamp, "Client Name: 0018"
]


def detect_margin(lines):
    """The dominant left-indent (in spaces) of lines that look like real tract/item
    headers -- used to tell a genuine item from a nested, indented sub-list."""
    from collections import Counter
    indents = []
    for line in lines:
        stripped = line.lstrip(' \t')
        indent = len(line) - len(stripped)
        if TRACT_PAT.match(stripped) or ITEM_PAT.match(stripped):
            indents.append(indent)
    if not indents:
        return 0
    return Counter(indents).most_common(1)[0][0]


PAGE_BREAK_MARKER = '\x0cPAGEBREAK\x0c'

# used only when the input is a .pdf (see pdf_to_indented_text): the left margin of a
# genuine tract/item line, and the point-to-space conversion for reconstructing indentation
PDF_BASE_X0 = 72.0
PDF_SPACE_WIDTH = 6.0


def pdf_to_indented_text(pdf_path):
    """Extract a PDF's text the way this script needs it: with each line's ORIGINAL
    left-indentation reconstructed from word positions (plain pdfplumber.extract_text()
    throws that away, which makes it impossible to tell a real item/tract line from a
    nested numbered sub-list inside a narrative -- e.g. an heirship affidavit's "1. ...",
    "2. ..." children), and with a page-break marker between pages so parse_document can
    tell a genuine multi-line footnote from one that's merely interrupted by a page turn.
    Requires pdfplumber (pip install pdfplumber)."""
    try:
        import pdfplumber
    except ImportError:
        print("This is a .pdf input, which requires the 'pdfplumber' package.")
        print("Install it with:  pip install pdfplumber")
        sys.exit(1)

    def cluster_lines(words, tol=3.0):
        lines, current, anchor = [], [], None
        for w in sorted(words, key=lambda w: (w['top'], w['x0'])):
            if anchor is None or abs(w['top'] - anchor) <= tol:
                current.append(w)
                anchor = w['top'] if anchor is None else (anchor * 0.7 + w['top'] * 0.3)
            else:
                lines.append(current)
                current = [w]
                anchor = w['top']
        if current:
            lines.append(current)
        return lines

    out_lines = []
    with pdfplumber.open(pdf_path) as pdf:
        for pi, page in enumerate(pdf.pages):
            if pi > 0:
                out_lines.append(PAGE_BREAK_MARKER)
            for line_words in cluster_lines(page.extract_words()):
                lw = sorted(line_words, key=lambda w: w['x0'])
                indent = max(0, round((lw[0]['x0'] - PDF_BASE_X0) / PDF_SPACE_WIDTH))
                out_lines.append(' ' * indent + ' '.join(w['text'] for w in lw))
    return '\n'.join(out_lines)


def find_boilerplate_lines(stripped_lines, num_pages):
    """A short line that repeats verbatim on (most of) EVERY page -- a running header like
    the client name, or a printed date -- is boilerplate, not content. This is deliberately
    tied to the page count, not a small fixed count: a document built from several near-
    duplicate deed chains (this one has three near-identical tracts) legitimately repeats
    the same clause 2-5 times, which a low fixed threshold would misfire on and silently
    delete from the middle of a real item. Only lines repeating on most pages qualify, and
    only when we actually know the page count (i.e. the input carries page-break markers);
    otherwise this returns nothing rather than guess."""
    if not num_pages or num_pages < 4:
        return set()
    from collections import Counter
    min_repeats = max(6, round(num_pages * 0.6))
    counts = Counter(s for s in stripped_lines if s and len(s) <= 100)
    return {s for s, c in counts.items() if c >= min_repeats}


def parse_document(raw_text):
    raw_lines = raw_text.split('\n')

    # page-break markers are structural, not content -- pull them out before anything else,
    # but remember which original-line index each one followed
    lines = []
    page_break_after = set()
    for line in raw_lines:
        if line.strip() == PAGE_BREAK_MARKER.strip():
            page_break_after.add(len(lines) - 1)
            continue
        lines.append(line)
    num_pages = len(page_break_after) + 1 if page_break_after else 0

    stripped_all = [l.strip() for l in lines]
    boilerplate = find_boilerplate_lines(stripped_all, num_pages)
    keep = [
        bool(s) and not any(p.match(s) for p in NOISE_PATS) and s not in boilerplate
        for s in stripped_all
    ]

    # margin detection needs the ORIGINAL (indented) lines, not the stripped text
    margin = detect_margin([line for line, k in zip(lines, keep) if k])
    margin_tol = 2  # a couple of stray spaces either side of the margin still counts

    kept_indices = [i for i, k in enumerate(keep) if k]
    content_lines = []  # (indent, text, page_break_follows)
    for pos, i in enumerate(kept_indices):
        line = lines[i]
        stripped = stripped_all[i]
        indent = len(line) - len(line.lstrip(' \t'))
        next_i = kept_indices[pos + 1] if pos + 1 < len(kept_indices) else len(lines)
        page_break_follows = any(j in page_break_after for j in range(i, next_i))
        content_lines.append((indent, stripped, page_break_follows))

    tracts = []
    footnotes = {}
    current_tract = None
    current_item_no = None
    current_item_lines = []
    current_footnote_no = None
    current_footnote_lines = []
    chain_counter = 0
    prev_item_no = None

    def flush_item():
        nonlocal current_item_no, current_item_lines
        if current_item_no is not None and current_tract is not None:
            text = re.sub(r'\s+', ' ', ' '.join(current_item_lines)).strip()
            current_tract['items'].append(
                {'item_no': current_item_no, 'text': text, 'chain': chain_counter}
            )
        current_item_no = None
        current_item_lines = []

    def flush_footnote():
        nonlocal current_footnote_no, current_footnote_lines
        if current_footnote_no is not None:
            text = re.sub(r'\s+', ' ', ' '.join(current_footnote_lines)).strip()
            footnotes[current_footnote_no] = text
        current_footnote_no = None
        current_footnote_lines = []

    for indent, line, page_break_follows in content_lines:
        on_margin = indent <= margin + margin_tol

        tm = TRACT_PAT.match(line) if on_margin else None
        if tm:
            flush_item()
            flush_footnote()
            current_tract = {'label': tm.group(1), 'name': tm.group(2), 'items': []}
            tracts.append(current_tract)
            chain_counter = 1
            prev_item_no = None
            continue

        im = ITEM_PAT.match(line) if on_margin else None
        if im:
            flush_item()
            flush_footnote()
            no = int(im.group(1))
            if prev_item_no is not None and no < prev_item_no:
                chain_counter += 1
            prev_item_no = no
            current_item_no = im.group(1)
            current_item_lines = [im.group(2)]
            continue

        fm = FOOTNOTE_PAT.match(line) if on_margin else None
        if fm:
            # a footnote is an ASIDE, not a new record -- pause (don't flush) whatever
            # item is in progress so its text can resume once the footnote closes
            flush_footnote()
            current_footnote_no = fm.group(1)
            current_footnote_lines = [fm.group(2).strip()]
            if page_break_follows:
                flush_footnote()
            continue

        # a plain continuation line: belongs to whichever block is currently open
        if current_footnote_no is not None:
            current_footnote_lines.append(line)
        elif current_item_no is not None:
            current_item_lines.append(line)
        # else: stray line before any tract/item/footnote (a title page line etc.) -- ignore

        if page_break_follows and current_footnote_no is not None:
            # footnotes in a title opinion never span a page break -- once the page turns,
            # resume appending to the item that was paused, not the footnote
            flush_footnote()

    flush_item()
    flush_footnote()
    return tracts, footnotes


# =======================================================================================
# STAGE 2 -- normalize each item's text into structured fields
# =======================================================================================

SENT_PAT = re.compile(
    r'^(?:As evidenced by|By)\s+(?P<instr_type>.+?)\s+dated\s+(?:effective\s+)?'
    r'(?P<date1>[A-Z][a-z]+ \d{1,2}, \d{4})'
    r'(?:, but (?:effective )?(?P<date2>[A-Z][a-z]+ \d{1,2}, \d{4}))?,?\s+'
    r'recorded in Volume (?P<vol>\d+), Page (?P<page>\d+),\s*'
    r'(?P<records_type>Deed Records|Official Public Records)'
    r'(?:\s*\(the "(?P<nickname>[^"]+)"\))?,\s*(?P<rest>.*)$'
)
# a correction-deed clause has its own internal "recorded in Volume/Page" and commas;
# strip the whole thing before SENT_PAT runs, keeping the ORIGINAL instrument's Volume/Page
CORRECTION_CLAUSE_RE = re.compile(
    r',\s*as corrected by .+?\s+recorded in Volume \d+, Page \d+, (?:Deed Records|Official Public Records)'
)
FOOTNOTE_MARK_PAT = re.compile(r'\.(\d{1,2})$')
# a footnote marker landing MID-sentence ("...1995.4 one-third...") shows up as a
# period+digit(s) immediately followed by a lowercase word -- never legitimate prose here
MIDTEXT_FOOTNOTE_PAT = re.compile(r'\.(\d{1,2})\s+(?=[a-z])')


def strip_footnote(text, footnotes):
    m_end = FOOTNOTE_MARK_PAT.search(text)
    if m_end and m_end.group(1) in footnotes:
        return m_end.group(1), text[:m_end.start()] + '.'
    m_mid = MIDTEXT_FOOTNOTE_PAT.search(text)
    if m_mid and m_mid.group(1) in footnotes:
        return m_mid.group(1), text[:m_mid.start()] + '.'
    return "", text


# interest is matched GREEDILY so the split lands on the LAST " to " in the sentence
GRANTOR_GRANTEE_PAT = re.compile(r'^(?P<grantor>.+?)\s+conveyed\s+(?P<interest>.+)\s+to\s+(?P<grantee>.+?)\.?\s*$')

# "X replaced Y (successor to Z) as Co-Trustee of the ... Trust ..." -- no "conveyed ... to"
# at all; the party stepping DOWN is the functional grantor, the one stepping IN is the
# grantee, named alongside the office/trust they now hold
REPLACED_TRUSTEE_PAT = re.compile(
    r'^(?P<new_party>.+?)\s+replaced\s+(?P<old_party>.+?)\s*(?:\([^)]*\)\s*)?'
    r'as\s+(?P<role>[A-Za-z][A-Za-z \-]*?Trustee)\s+of\s+(?P<trust>.+?)'
    r'(?:,?\s+as amended)?(?:,\s+on\s+[A-Z][a-z]+ \d{1,2}, \d{4})?\.?\s*$'
)

# "X and Y partitioned amongst themselves ..." -- the same parties are both grantor and
# grantee (no interest changes hands with anyone outside the group)
SELF_PARTITION_PAT = re.compile(r'^(?P<parties>.+?)\s+partitioned amongst themselves\b')

# "X, Y and Z stipulated to the ownership of ..." -- names the parties agreeing, but
# (unlike a conveyance) never names a distinct recipient
STIPULATION_NO_GRANTEE_PAT = re.compile(r'^(?P<parties>.+?)\s+stipulated to the ownership\b')

# the opening line of a multi-generation heirship affidavit ("a. H. M. Roberts died
# testate...", "(i) Eulalia Sanders Blackmon died intestate...") -- used only as a last
# resort, to name a single representative decedent when the narrative names several
HEIRSHIP_FIRST_DECEDENT_RE = re.compile(
    r'(?:^|[:.]\s*)(?:[a-z]\.|\([ivx]+\))\s*(?P<name>[A-Z][A-Za-z.\']+(?:\s+[A-Z][A-Za-z.\']+){0,4})'
    r'\s+died\s+(?:testate|intestate)\b'
)

SHORT_SUFFIXES = {'jr', 'jr.', 'sr', 'sr.', 'ii', 'iii', 'iv', 'v', 'trustee', 'et ux', 'et ux.', 'et vir', 'et vir.',
                   'deceased', 'inc', 'inc.', 'llc', 'l.l.c.', 'ltd', 'ltd.', 'lp', 'l.p.', 'co', 'co.', 'corp',
                   'corp.', 'p.c.', 'pc', 'n.a.', 'na'}
DESCRIPTOR_PREFIXES = ('a widow', 'widow', 'a widower', 'widower', 'a single', 'single man', 'single woman',
                        'individually', 'in her', 'in his', 'as trustee', 'as independent', 'as executor',
                        'as executrix', 'as administrator', 'as guardian', 'a married', 'as co-trustee',
                        'as successor')
# an alias ("a/k/a Other Name") is dropped entirely -- only the primary name is kept
AKA_TOKEN_RE = re.compile(r'^(?:a/?k/?a\b)', re.I)
AKA_INLINE_RE = re.compile(r',\s*a/?k/?a\s+[^,]+', re.I)
# a (NN%) share annotation attached to a party -- stripped globally
PERCENT_PAREN_RE = re.compile(r'\s*\(\d[\d.]*%\)')
# transactional qualifiers describing the CONVEYANCE, not the person/entity -- dropped
TRAILING_QUALIFIER_RE = re.compile(
    r',?\s*as\s+(?:(?:his|her|their|a)\s+)?(?:sole and )?separate property\.?$'
    r'|,?\s*dealing in (?:his|her|their) separate property\.?$'
    r'|,?\s*in equal shares\.?$',
    re.I,
)
DROP_TOKEN_RE = re.compile(
    r'^as\s+(?:(?:his|her|their|a)\s+)?(?:sole and )?separate property$'
    r'|^dealing in (?:his|her|their) separate property$'
    r'|^in equal shares$|^individually$',
    re.I,
)


def strip_qualifiers(s):
    return TRAILING_QUALIFIER_RE.sub('', s).strip()


# "X and wife, Y" / "X joined by husband, Y" -- the connector is dropped and X, Y become
# two SEPARATE lines
SPOUSE_END_RE = re.compile(r'\b(?:and|joined by) (his wife|her husband|wife|husband)$', re.I)
# a date/deed/trust description inside a party string -- splitting on commas would shred
# the clause, so leave it untouched
NON_SPLITTABLE_RE = re.compile(
    r'\d|dated|recorded|agreement|trust under|which |correct|between|as amended|revocable trust|living trust',
    re.I,
)


def split_names(s):
    """Turn a running prose list of grantor/grantee parties into one name per line,
    keeping each party (name + spouse/suffix/descriptor) together on a single line,
    but splitting an "X and wife, Y" couple into two separate lines."""
    if not s:
        return s
    s = PERCENT_PAREN_RE.sub('', s)
    s = AKA_INLINE_RE.sub('', s)
    s = s.strip().rstrip(' ,.;')
    if not s:
        return s
    if NON_SPLITTABLE_RE.search(s):
        return strip_qualifiers(s)
    tokens = [t.strip() for t in s.split(',') if t.strip()]
    merged = []
    pending_spouse_next = False

    def emit(text):
        sm = SPOUSE_END_RE.search(text)
        if sm:
            merged.append(text[:sm.start()].strip())
            return True
        merged.append(text)
        return False

    for tok in tokens:
        low = tok.lower()
        if pending_spouse_next:
            merged.append(tok)
            pending_spouse_next = False
            continue
        if DROP_TOKEN_RE.match(tok) or AKA_TOKEN_RE.match(tok):
            continue
        if low.startswith('and '):
            tok_clean = tok[4:].strip()
            if tok_clean.lower() in ('wife', 'husband', 'his wife', 'her husband'):
                pending_spouse_next = True
                continue
            pending_spouse_next = emit(tok_clean)
            continue
        is_short_suffix = low in SHORT_SUFFIXES
        is_descriptor = any(low.startswith(p) for p in DESCRIPTOR_PREFIXES)
        if merged and (is_short_suffix or is_descriptor):
            merged[-1] = merged[-1] + ', ' + tok
            continue
        pending_spouse_next = emit(tok)
    merged = [strip_qualifiers(x) for x in merged]
    merged = [x for x in merged if x]
    expanded = []
    for line in merged:
        expanded.extend(maybe_split_bare_and(line))
    return '\n'.join(expanded)


COMPANY_HINT_RE = re.compile(
    r'\b(Company|Co\.|LLC|L\.L\.C\.|Inc\.|Ltd\.|LP|L\.P\.|Trust|Trustee|Corporation|Corp\.|Partners|Partnership|Bank)\b',
    re.I,
)


def maybe_split_bare_and(line):
    """A list without an Oxford comma ("A, B, C and D") leaves the last two names in
    one comma-token -- split that into two people, unless it's a single company/trust
    name or an "X and wife Y" couple (kept together upstream already)."""
    if ' and ' not in line.lower():
        return [line]
    if COMPANY_HINT_RE.search(line) or re.search(r'\band (wife|husband)\b', line, re.I):
        return [line]
    parts = re.split(r'\s+and\s+', line)
    if len(parts) == 2 and all(re.match(r'^[A-Z]', p.strip()) for p in parts):
        return [p.strip() for p in parts]
    return [line]


CUT_MARKERS = ['. "', '. This ', '. Furthermore', '. See Requirement', '. See the ', '. We have', '. Opinion']
# a period+digit is usually a new sentence ("...Allred. 2. ...") but NOT when the period
# closes an abbreviation like "No." ("RD No. 1 to CrownRock...", "Cause No. 725...")
DIGIT_CLAUSE_RE = re.compile(r'(?<![Nn]o)\.\s+\d')


def first_sentence(rest):
    """The real grantor/interest/grantee only ever live in the FIRST sentence of the
    clause following the header -- cut off any supplementary sentence, quoted excerpt,
    citation, or semicolon-separated proviso that follows it."""
    cut_at = len(rest)
    for marker in CUT_MARKERS:
        idx = rest.find(marker)
        if idx != -1:
            cut_at = min(cut_at, idx + 1)
    dm = DIGIT_CLAUSE_RE.search(rest)
    if dm:
        cut_at = min(cut_at, dm.start() + 1)
    semi_idx = rest.find(';')
    if semi_idx != -1:
        cut_at = min(cut_at, semi_idx + 1)
    return rest[:cut_at]


# "...conveyed X to GRANTEE, and reserved [interest] ..." -- describes what the GRANTOR
# kept, not who the grantee is
RESERVATION_CLAUSE_RE = re.compile(r',?\s+and reserved\b.*$', re.I)

FOLLOWING_PARTIES_PAT = re.compile(
    r'^(?P<pre>.+?)\s+to the following parties,\s*in the proportions set forth opposite their names:\s*'
    r'(?P<list>.+?)\.?\s*$'
)
NUM_SPLIT_PAT = re.compile(r'\s*\d[\d,]*(?:\.\d+)?%\s*|\s*\d[\d,]*/\d[\d,]*\s*')


def cap_line(s):
    s = s.strip()
    if s and s[0].islower():
        return s[0].upper() + s[1:]
    return s


def split_proportion_list(list_text):
    """"Name1 frac1 Name2 frac2 ..." (a distribution table flattened to one line) ->
    just the names, one per line, with the proportion numbers dropped entirely."""
    names = [n.strip().strip(',').strip() for n in NUM_SPLIT_PAT.split(list_text)]
    names = [n for n in names if n]
    return '\n'.join(cap_line(n) for n in names)


# ---- probate / "died testate" items -------------------------------------------------
PROBATE_INTRO_PAT = re.compile(
    r'^(?P<decedent>.+?) died (?:testate|intestate) on (?P<death_date>[A-Z][a-z]+ \d{1,2}, \d{4})\.\s*(?P<remainder>.*)$',
    re.S,
)
PROBATE_ORDER_PAT = re.compile(
    r'admitted to probate(?: as a muniment of title)? by Order dated (?P<order_date>[A-Z][a-z]+ \d{1,2}, \d{4}),'
    r' under Cause No\.\s*[\d,]+,\s*[^,]+,\s*[^,]+ County, Texas'
    r'(?:, certified copies of which are recorded in Volume (?P<vol>\d+), Page (?P<page>\d+), (?P<records_type>Deed Records))?',
    re.S,
)
PROBATE_WILL_PAT = re.compile(r'In (?:his|her) Will, .+? devised (?P<bequest>.+)$', re.S)
BEQUEST_FOLLOWING_PARTIES_PAT = re.compile(
    r'^(?P<pre>.+?)\s+to the following parties,\s*in the proportions set forth opposite their names:\s*'
    r'(?P<list>.+?)\.?\s*(?:Furthermore.*)?$',
    re.S,
)
BEQUEST_EQUAL_SHARES_PAT = re.compile(
    r'^(?P<pre>.+?)\s+to (?:his|her|their) .*?,\s*(?P<names>.+?),?\s*in equal shares\.?\s*$', re.S)
BEQUEST_SINGLE_PAT = re.compile(
    r'^(?P<pre>.+?)\s+to (?:his|her) (?:daughter|son|wife|husband|sole heir),\s*(?P<name>[^.]+?)\.\s*$', re.S)


def extract_probate(text):
    """Handles the "[Decedent] died testate/intestate on [date]. ... admitted to
    probate by Order dated [date] ... In [his/her] Will, [Decedent] devised ... to
    [beneficiaries]." pattern -- a wholly different shape from the main SENT_PAT."""
    intro = PROBATE_INTRO_PAT.match(text)
    if not intro:
        return None
    decedent = intro.group('decedent').strip()
    remainder = intro.group('remainder')

    order_m = PROBATE_ORDER_PAT.search(remainder)
    order_date = vol = page = records_type = ''
    if order_m:
        order_date = order_m.group('order_date')
        vol = order_m.group('vol') or ''
        page = order_m.group('page') or ''
        records_type = order_m.group('records_type') or 'Deed Records'

    will_m = PROBATE_WILL_PAT.search(remainder)
    if not will_m:
        return None
    bequest = will_m.group('bequest').strip()

    interest = beneficiaries = ''
    fp = BEQUEST_FOLLOWING_PARTIES_PAT.match(bequest)
    if fp:
        interest = fp.group('pre').strip().rstrip(',')
        beneficiaries = split_proportion_list(fp.group('list'))
    else:
        eq = BEQUEST_EQUAL_SHARES_PAT.match(bequest)
        if eq:
            interest = eq.group('pre').strip().rstrip(',')
            beneficiaries = split_names(eq.group('names'))
        else:
            single = BEQUEST_SINGLE_PAT.match(bequest)
            if single:
                interest = single.group('pre').strip().rstrip(',')
                beneficiaries = cap_line(strip_qualifiers(single.group('name').strip()))

    if not beneficiaries:
        return None

    return {
        'grantor': f"{decedent}\nEstate of {decedent}, Dec'd",
        'order_date': order_date,
        'vol': vol, 'page': page,
        'records_type': records_type or 'Deed Records',
        'interest': interest,
        'grantee': beneficiaries,
    }


def parse_date(s):
    if not s:
        return ""
    try:
        return datetime.strptime(s, "%B %d, %Y").date()
    except ValueError:
        return s


def build_records(tracts, footnotes):
    rows = []
    for t in tracts:
        for it in t['items']:
            text = it['text']
            footnote_ref, text_wo_fn = strip_footnote(text, footnotes)
            text_wo_fn = CORRECTION_CLAUSE_RE.sub('', text_wo_fn)

            rec = {
                'tract_label': t['label'], 'tract_name': t['name'], 'chain': it['chain'],
                'item_no': int(it['item_no']),
                'instr_type': '', 'date_executed': '', 'date_effective': '',
                'volume': '', 'page': '', 'records_type': '', 'nickname': '',
                'grantor': '', 'interest': '', 'grantee': '',
                'footnote_ref': footnote_ref, 'footnote_text': footnotes.get(footnote_ref, ''),
                'full_text': text,
            }

            m = SENT_PAT.match(text_wo_fn)
            if m:
                rec['instr_type'] = m.group('instr_type').strip()
                rec['date_executed'] = parse_date(m.group('date1'))
                rec['date_effective'] = parse_date(m.group('date2')) if m.group('date2') else ''
                rec['volume'] = m.group('vol')
                rec['page'] = m.group('page')
                rec['records_type'] = m.group('records_type')
                rec['nickname'] = m.group('nickname') or ''
                rest = first_sentence(m.group('rest'))
                fp = FOLLOWING_PARTIES_PAT.match(rest)
                if fp:
                    pre = fp.group('pre')
                    gm_pre = re.match(r'^(?P<grantor>.+?)\s+conveyed\s+(?P<interest>.+)$', pre)
                    if gm_pre:
                        grantor_lines = split_names(gm_pre.group('grantor'))
                        rec['grantor'] = '\n'.join(cap_line(x) for x in grantor_lines.split('\n'))
                        rec['interest'] = gm_pre.group('interest').strip()
                    else:
                        rec['interest'] = pre.strip()
                    rec['grantee'] = split_proportion_list(fp.group('list'))
                else:
                    gm = GRANTOR_GRANTEE_PAT.match(rest)
                    rp = REPLACED_TRUSTEE_PAT.match(rest) if not gm else None
                    sp = SELF_PARTITION_PAT.match(rest) if not gm and not rp else None
                    sn = STIPULATION_NO_GRANTEE_PAT.match(rest) if not gm and not rp and not sp else None
                    if gm:
                        grantor_lines = split_names(RESERVATION_CLAUSE_RE.sub('', gm.group('grantor')))
                        rec['grantor'] = '\n'.join(cap_line(x) for x in grantor_lines.split('\n'))
                        rec['interest'] = gm.group('interest').strip()
                        grantee_lines = split_names(RESERVATION_CLAUSE_RE.sub('', gm.group('grantee')))
                        rec['grantee'] = '\n'.join(cap_line(x) for x in grantee_lines.split('\n'))
                    elif rp:
                        # "X replaced Y ... as Co-Trustee of Z" -- no conveyance verb at all
                        rec['grantor'] = cap_line(rp.group('old_party').strip().rstrip(','))
                        rec['grantee'] = f"{cap_line(rp.group('new_party').strip())}, {rp.group('role').strip()} of {rp.group('trust').strip().rstrip(',')}"
                        rec['interest'] = rest.strip()
                    elif sp:
                        # "X and Y partitioned amongst themselves ..." -- same people on both sides
                        parties = split_names(sp.group('parties'))
                        parties_fmt = '\n'.join(cap_line(x) for x in parties.split('\n'))
                        rec['grantor'] = parties_fmt
                        rec['grantee'] = parties_fmt
                        rec['interest'] = rest.strip()
                    elif sn:
                        # "X, Y and Z stipulated to the ownership of ..." -- no recipient is named
                        parties = split_names(sn.group('parties'))
                        rec['grantor'] = '\n'.join(cap_line(x) for x in parties.split('\n'))
                        rec['interest'] = rest.strip()
                    else:
                        rec['interest'] = rest.strip()
            else:
                probate = extract_probate(text_wo_fn)
                if probate:
                    rec['instr_type'] = 'Probate'
                    rec['date_executed'] = parse_date(probate['order_date'])
                    rec['volume'] = probate['vol']
                    rec['page'] = probate['page']
                    rec['records_type'] = probate['records_type']
                    rec['grantor'] = probate['grantor']
                    rec['interest'] = probate['interest']
                    rec['grantee'] = probate['grantee']

            if not rec['grantor'] or not rec['grantee']:
                # last resort: a multi-generation heirship narrative that named several
                # decedents -- pick the FIRST one named as a single representative party
                hm = HEIRSHIP_FIRST_DECEDENT_RE.search(text_wo_fn)
                if hm and not rec['grantor'] and not rec['grantee']:
                    name = re.sub(r'\s+', ' ', hm.group('name')).strip().rstrip('.')
                    rec['grantor'] = f"Estate of {name}, Dec'd"
                    rec['grantee'] = f"Heirs of {name}, see instrument"

            # absolute last resort: never leave a row with no grantor/grantee at all --
            # flag it so it can be found and reviewed instead of silently disappearing
            if not rec['grantor']:
                rec['grantor'] = 'Unclear, see instrument'
            if not rec['grantee']:
                rec['grantee'] = 'Unclear, see instrument'

            rows.append(rec)
    return rows


# =======================================================================================
# STAGE 3 -- classify each item and emit the normalized CD / NCD land-software import files
#
# This mirrors, field-for-field, the classification and output conventions of the user's
# existing runsheet normalizer (target_cols schema, SUBSTRING_RULES / EXACT_RULES,
# assign_subtype_and_type, assign_booktype, Notes concatenation, CD/NCD split, styling).
# =======================================================================================

STATE_ABBR = {
    'TEXAS': 'TX', 'NEW MEXICO': 'NM', 'OKLAHOMA': 'OK', 'LOUISIANA': 'LA',
    'COLORADO': 'CO', 'KANSAS': 'KS', 'NORTH DAKOTA': 'ND', 'WYOMING': 'WY',
    'MONTANA': 'MT', 'UTAH': 'UT', 'ARKANSAS': 'AR', 'MISSISSIPPI': 'MS',
    'CALIFORNIA': 'CA', 'ALABAMA': 'AL', 'PENNSYLVANIA': 'PA', 'OHIO': 'OH',
}
KNOWN_ABBRS = set(STATE_ABBR.values())

# High-confidence, property-specific phrasing: "situated in X County, Y", "lands located
# in X County, Y", "Deed Records of X County, Y" / "Official Public Records of X County, Y".
PROPERTY_HINT_RE = re.compile(
    r'(?:situated|located)\s+in\s+(?P<c1>[A-Z][A-Za-z.\' ]{2,30}?)\s+County,\s*(?P<s1>[A-Za-z]{2,20})'
    r'|(?:Deed Records|Official Public Records)\s+of\s+(?P<c2>[A-Z][A-Za-z.\' ]{2,30}?)\s+County,\s*(?P<s2>[A-Za-z]{2,20})'
)

# Any "X County, Y" mention -- used only as a fallback, and only when it does NOT sit near
# court/probate venue language (those name where a decedent lived or where a case was filed,
# not where the property is).
COUNTY_STATE_RE = re.compile(r'([A-Z][A-Za-z.\' ]{2,30}?)\s+County,\s*([A-Za-z]{2,20})\b')
COURT_CONTEXT_RE = re.compile(
    r'Cause No\.|County Court|Probate Court|Judicial District Court|died|survived by|Clerk of', re.I
)


def normalize_state(raw):
    s = raw.strip().upper()
    if s in KNOWN_ABBRS:
        return s
    return STATE_ABBR.get(s, raw.strip()[:2].upper())


def detect_county_state(raw_text):
    """Best-effort detection of the property's county/state from the document body itself
    (title-opinion narratives rarely have a dedicated header field for it). Returns
    (county, state, how) -- how explains where the value came from, for the run summary."""
    hints = Counter()
    for m in PROPERTY_HINT_RE.finditer(raw_text):
        county = (m.group('c1') or m.group('c2')).strip()
        state = normalize_state(m.group('s1') or m.group('s2'))
        hints[(county, state)] += 1
    if hints:
        (county, state), _ = hints.most_common(1)[0]
        return county, state, 'detected from phrases like "located/situated in ... County" or "Deed Records of ... County"'

    counts = Counter()
    for m in COUNTY_STATE_RE.finditer(raw_text):
        context = raw_text[max(0, m.start() - 60):m.start()]
        if COURT_CONTEXT_RE.search(context):
            continue
        counts[(m.group(1).strip(), normalize_state(m.group(2)))] += 1
    if counts:
        (county, state), _ = counts.most_common(1)[0]
        return county, state, 'detected by frequency of "... County, State" mentions (excluding court/probate venues)'

    return None, None, 'not found in document text'


TARGET_COLS = [
    "Class", "Type", "Subtype", "Grantor", "Grantee", "Date", "Recorded",
    "Inst.Date", "Acknowledged", "Filed", "Booktype", "Book", "Page",
    "Inst.No.", "County", "State", "Notes", "Requirements", "Essences",
    "Flags", "Restrictions", "Warnings", "Reviews", "Files",
]

SUBSTRING_RULES = {
    "ASSIGNMENT OF OIL AND GAS LEASE": "AOGL",
    "PARTIAL ASSIGNMENT OF OIL AND GAS LEASE": "PART AOGL",
    "ASSIGNMENT OF OIL, GAS, AND MINERAL LEASE": "AOGML",
    "PARTIAL ASSIGNMENT OF OIL, GAS, AND MINERAL LEASE": "PART AOGML",
    "ASSIGNMENT OF OVERRIDING ROYALTY INTEREST": "AORRI",
    "ASSIGNMENT OF OVERRIDING ROAYLTY INTEREST": "AORRI",
    "ASSIGNMENT OF DEED OF TRUST": "ASSIGN DT",
    "ASSIGNMENT OF PRODUCTION": "DT",
    "ASSIGNMENT OF AS-EXTRACTED COLLATERAL": "DT",
    "BILL OF SALE AND ASSIGNMENT": "ABOS",
    "ASSIGNMENT AND BILL OF SALE": "ABOS",
    "PARTIAL ASSIGNMENT": "PART ASSIGN",
    "ASSIGNMENT": "ASSIGN",
    "BLANKET MINERAL AND ROYALTY DEED": "MD",
    "MINERAL AND ROYALTY DEED": "MRD",
    "MINERAL/ROYALTY DEED": "MRD",
    "MINERAL AND ROAYLTY INTEREST DEED": "MRD",
    "CORRECTION MINERAL AND ROYALTY DEED": "MRD",
    "MINERAL DEED": "MD",
    "INDEPENDENT CO-EXECUTOR'S DISTRIBUTION DEED": "DIST DEED",
    "DISTRIBUTION DEED": "DIST DEED",
    "TRUSTEE'S DISTRIBUTION DEED": "DIST DEED",
    "DEED OF DISTRIBUTION": "DIST DEED",
    "AFFIDAVIT OF HEIRSHIP": "AFFT HEIR",
    "AFF OF HEIRSHIP": "AFFT HEIR",
    "AFFIDAVIT OF SUCCESSOR TRUSTEE": "AFFT SUCCESSOR TRST",
    "AFFIDAVIT": "AFFT",
    "SURFACE USE AND EASEMENT AGREEMENT": "AGREE",
    "EASEMENT": "EAS",
    "DEED OF TRUST": "DT",
    "QUIT CLAIM": "QCD",
    "GENERAL WARRANTY DEED": "GWD",
    "SPECIAL WARRANTY DEED": "SWD",
    "WARRANTY DEED WITH VENDOR'S LIEN": "WDVL",
    "WARRANTY DEED": "WD",
    "GIFT DEED": "GIFT DEED",
    "ADMINISTRATOR'S DEED": "ADMIN DEED",
    "EXECUTOR'S DEED": "EXEC DEED",
    "PARTITION DEED": "PART DEED",
    "PERSONAL REPRESENTATIVE'S DEED": "PR DEED",
    "RIGHT OF WAY DEED": "ROW DEED",
    "SHERIFF'S DEED": "SHRF DEED",
    "SUCCESSOR TRUSTEE'S DEED": "SUCCESSOR TRST DEED",
    "TRUSTEE'S DEED": "TRST DEED",
    "ROYALTY DEED": "RD",
    "DEED": "DEED",
    "BILL OF SALE": "BOS",
    "PARTIAL RELEASE OF LIEN": "RELL OF LIEN",
    "RELEASE OF LIEN": "RELL OF LIEN",
    "LIEN": "RELL OF LIEN",
    "RELEASE": "REL",
    "AMENDED AND RESTATED": "AMEND",
    "AMEND": "AMEND",
    "MEMORANDUM OF OIL AND GAS LEASE": "MEMO OGL",
    "MEMORANDUM OF OIL, GAS, AND MINERAL LEASE": "MEMO OGML",
    "OIL AND GAS LEASE": "OGL",
    "OIL & GAS LEASE": "OGL",
    "OIL, GAS, AND MINERAL LEASE": "OGML",
    "CC PROBATE": "PROBATE",
    "CC WILL": "PROBATE",
    "CC WILL / ORDER": "PROBATE",
    "WILL & PROBATE": "PROBATE",
    "CC/PROBATE": "PROBATE",
    "PROBATE": "PROBATE",
    "ESTATE": "PROBATE",
    "AGREEMENT": "AGREE",
    "CONVEYANCE": "CONV",
}
SORTED_SUBSTRING_RULES = sorted(SUBSTRING_RULES.items(), key=lambda x: len(x[0]), reverse=True)

EXACT_RULES = {
    r'\bAOGL\b': 'AOGL', r'\bAOGML\b': 'AOGML', r'\bAORRI\b': 'AORRI', r'\bASORI\b': 'AORRI',
    r'\bASSIGN DT\b': 'ASSIGN DT', r'\bABOS\b': 'ABOS', r'\bASSN\b': 'ASSIGN', r'\bASGN\b': 'ASSIGN',
    r'\bASN\b': 'ASSIGN', r'^AS$': 'ASSIGN', r'\bPART AOGL\b': 'PART AOGL', r'\bPART AOGML\b': 'PART AOGML',
    r'\bPART ASSIGN\b': 'PART ASSIGN', r'\bPART ASGN\b': 'PART ASSIGN', r'\bMD\b': 'MD', r'\bMRD\b': 'MRD',
    r'\bDIST DEED\b': 'DIST DEED', r'\bAFFT HEIR\b': 'AFFT HEIR', r'\bAFF H\b': 'AFFT HEIR', r'\bAFFT\b': 'AFFT',
    r'\bAGREE\b': 'AGREE', r'\bAGMT\b': 'AGREE', r'\bDOTO\b': 'DT', r'\bD OF TR\b': 'DT', r'\bDT\b': 'DT',
    r'\bQCD\b': 'QCD', r'\bWD\b': 'WD', r'\bCOR WD\b': 'WD', r'\bCORRECTION WD\b': 'WD', r'\bWD RECORD\b': 'WD', r'\bWD/ VL\b': 'WD',
    r'\bGWD\b': 'GWD', r'\bSWD\b': 'SWD', r'\bWDVL\b': 'WDVL', r'\bBOS\b': 'BOS', r'\bCONV\b': 'CONV',
    r'\bREL\b': 'REL', r'\bREL LN\b': 'RELL OF LIEN', r'\bREL. OL\b': 'REL',
    r'\bAMD\b': 'AMEND', r'\bOGL\b': 'OGL', r'\bOGML\b': 'OGML', r'\bPROB\b': 'PROBATE', r'\bCC PROB\b': 'PROBATE',
    r'\bWILL\b': 'PROBATE', r'\bPATENT\b': 'PATENT', r'\bHEIRSHIP\b': 'AFFT HEIR', r'\bO&GL\b': 'OGL', r'\bOG&ML\b': 'OGML', r'\bMOGL\b': 'OGL',
    r'\bJDG\b': 'JUDGMENT',
}


def assign_subtype_and_type(instrument):
    """Ported verbatim (rule tables + logic) from the user's existing runsheet normalizer."""
    if not instrument:
        return "MISC", "MISC.", "NCD"

    instr = str(instrument).upper().strip()
    subtype = ""

    for phrase, sub_val in SORTED_SUBSTRING_RULES:
        if phrase in instr:
            subtype = sub_val
            break

    if not subtype:
        for pattern, sub_val in EXACT_RULES.items():
            if re.search(pattern, instr):
                subtype = sub_val
                break

    if not subtype:
        subtype = instr.split('\n')[0].strip()

    conveyances = ["MD", "WD", "GWD", "SWD", "QCD", "DEED", "CONV", "BOS", "PROBATE", "AFFT HEIR", "MRD", "DIST DEED", "GIFT DEED", "ADMIN DEED", "EXEC DEED", "PART DEED", "PR DEED", "SHRF DEED", "SUCCESSOR TRST DEED", "TRST DEED", "RD", "ROW DEED", "WDVL", "PATENT"]
    assignments = ["ASSIGN", "ABOS", "AOGL", "AOGML", "PART ASSIGN", "PART AOGL", "PART AOGML", "AORRI", "ASSIGN DT", "BLM ASSIGN", "BLM TRANSFER", "TOD"]
    leases = ["OGL", "OGML"]
    releases = ["REL"]

    if any(c == subtype or f" {c} " in f" {subtype} " for c in conveyances) or "PATENT" in instr:
        return subtype, "Conveyance", "CD"
    if any(a == subtype or f" {a} " in f" {subtype} " for a in assignments):
        return subtype, "Assignment", "CD"
    if any(l == subtype or f" {l} " in f" {subtype} " for l in leases):
        return subtype, "Lease", "CD"
    if any(r == subtype or f" {r} " in f" {subtype} " for r in releases):
        return subtype, "Release", "CD"

    doc_class = "NCD"
    easements = ["EAS", "ROW", "ROW/EAS", "EASEMENT"]
    d_of_trust = ["DT", "MTG", "SUB", "DEED OF TRUST", "D OF TR"]
    liens = ["MML", "FTL", "STL", "HL", "AJ", "RELL OF LIEN", "LIEN"]
    agreements = ["AGREE", "TRUST AGREE", "STIP", "JOA", "RDO", "AGREEMENT", "CONTRACT"]
    judgements = ["JUDG", "JUDGMENT"]

    if any(d == subtype or f" {d} " in f" {subtype} " for d in d_of_trust) or "DEED OF TRUST" in instr or "D OF TR" in instr or re.search(r'\bDT\b', instr):
        return "DT", "DEED OF TRUST", doc_class
    elif any(e == subtype or f" {e} " in f" {subtype} " for e in easements) or "EASEMENT" in instr:
        return "Easement", "Easement", doc_class
    elif any(j == subtype or f" {j} " in f" {subtype} " for j in judgements) or "JUDG" in instr:
        return "JUD", "MISC.", doc_class
    elif "AMEND" in subtype or "AMEND" in instr:
        return "AMENDM", "MISC.", doc_class
    elif any(a == subtype or f" {a} " in f" {subtype} " for a in agreements) or "AGREEMENT" in instr or "CONTRACT" in instr:
        return "AGREMNT", "AGRE", doc_class
    elif any(l == subtype or f" {l} " in f" {subtype} " for l in liens) or "LIEN" in instr:
        return "RELL OF LIEN", "Lien", doc_class
    else:
        return "MISC", "MISC.", doc_class


def assign_booktype(records_type, inst_date_str):
    """Records Type (Deed Records / Official Public Records) takes priority, same as the
    existing normalizer; falls back to a year cutoff (<1990 -> DR, >=1990 -> OPR)."""
    if records_type:
        rt = records_type.strip().upper()
        if 'DEED' in rt:
            return 'DR'
        if 'OFFICIAL' in rt or 'OPR' in rt or 'PUBLIC' in rt:
            return 'OPR'
    m = re.search(r'\d{4}', inst_date_str or '')
    if m:
        return "DR" if int(m.group()) < 1990 else "OPR"
    return ""


def fmt_date(d):
    """Format a date object (or already-formatted string) as MM/DD/YYYY, matching the
    existing normalizer's clean_specific_dates() output -- built directly here (never via
    str()-then-reparse) so real date objects can't get corrupted into ambiguous strings."""
    if not d:
        return ""
    if isinstance(d, date):
        return d.strftime('%m/%d/%Y')
    return str(d)


def build_notes(rec):
    """Concatenate every field that has no dedicated target_cols slot, mirroring the
    existing normalizer's "ColumnName: value" Notes convention, plus the auto-flagged
    Judgment/Agreement notes it appends based on the instrument text."""
    parts = []
    if rec.get('tract_label'):
        parts.append(f"Tract: {rec['tract_label']}")
    if rec.get('tract_name'):
        parts.append(f"Tract Description: {rec['tract_name']}")
    if rec.get('chain'):
        parts.append(f"Chain: {rec['chain']}")
    if rec.get('item_no'):
        parts.append(f"Item #: {rec['item_no']}")
    if rec.get('date_effective'):
        parts.append(f"Effective Date: {fmt_date(rec['date_effective'])}")
    if rec.get('nickname'):
        parts.append(f"Deed Nickname: {rec['nickname']}")
    if rec.get('interest'):
        parts.append(f"Interest / Conveyance: {rec['interest']}")
    if rec.get('footnote_ref'):
        parts.append(f"Footnote #: {rec['footnote_ref']}")
    if rec.get('footnote_text'):
        parts.append(f"Footnote Text: {rec['footnote_text']}")
    if rec.get('full_text'):
        parts.append(f"Full Item Text: {rec['full_text']}")

    instr_val = rec.get('instr_type', '')
    if instr_val:
        parts.append(f"Original Instrument: {instr_val}")
        instr_upper = instr_val.upper()
        if len(instr_val) >= 75 or "CONSOLIDATED DT" in instr_upper or "SUPPLEMENT TO DT" in instr_upper:
            parts.append("Check instrument couldnt determine what it is form the runsheet import")
        if any(j in instr_upper for j in ["JUDG", "JUDGMENT"]):
            parts.append("Note: This document is a Judgment.")
        if any(a in instr_upper for a in ["AGREE", "TRUST AGREE", "STIP", "JOA", "CONTRACT"]):
            parts.append("Note: This document is an Agreement.")

    return "\n\n".join(parts)


def build_normalized_rows(rows, county, state):
    """Map the parsed History-of-Title rows into the target_cols schema used by the
    land-software import (same schema/classification as the existing runsheet normalizer)."""
    out = []
    for rec in rows:
        subtype, doc_type, doc_class = assign_subtype_and_type(rec.get('instr_type', ''))
        inst_date_str = fmt_date(rec.get('date_executed', ''))
        out.append({
            "Class": doc_class,
            "Type": doc_type,
            "Subtype": subtype,
            "Grantor": rec.get('grantor', ''),
            "Grantee": rec.get('grantee', ''),
            "Date": "",
            "Recorded": "",
            "Inst.Date": inst_date_str,
            "Acknowledged": "",
            "Filed": "",
            "Booktype": assign_booktype(rec.get('records_type', ''), inst_date_str),
            "Book": rec.get('volume', ''),
            "Page": rec.get('page', ''),
            "Inst.No.": "",
            "County": county,
            "State": state,
            "Notes": build_notes(rec),
            "Requirements": "",
            "Essences": "",
            "Flags": "",
            "Restrictions": "",
            "Warnings": "",
            "Reviews": "",
            "Files": "",
        })
    return out


def save_normalized_excel(records, out_path):
    """Same header fill / font / border / wrap conventions as the existing normalizer's
    save_formatted_excel()."""
    if not records:
        print(f"No records to write to {out_path}.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Normalized Runsheet"

    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(TARGET_COLS)
    for c_idx in range(1, len(TARGET_COLS) + 1):
        cell = ws.cell(row=1, column=c_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for rec in records:
        ws.append([rec.get(col, '') for col in TARGET_COLS])

    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(TARGET_COLS)):
        for cell in row_cells:
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    widths = {'A': 10, 'B': 15, 'C': 20, 'D': 25, 'E': 25, 'H': 15, 'J': 15, 'N': 15, 'O': 15, 'Q': 50, 'S': 50}
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

    wb.save(out_path)
    print(f"File generated: {out_path} ({len(records)} records)")


# =======================================================================================
# Companion outputs: the original "History of Title" master sheet (17 columns, one row
# per item -- feed this into the existing runsheet normalizer to get its own CD/NCD split,
# to compare against this script's own CD/NCD output for the same document) and a plain-
# text report of every field this script could not fill in with confidence.
# =======================================================================================

HOT_HEADERS = ["Tract", "Tract Description", "Chain", "Item #", "Instrument Type",
               "Date Executed", "Effective Date", "Volume", "Page", "Records Type",
               "Deed Nickname", "Grantor", "Interest / Conveyance", "Grantee",
               "Footnote #", "Footnote Text", "Full Item Text"]
HOT_KEYS = ["tract_label", "tract_name", "chain", "item_no", "instr_type",
            "date_executed", "date_effective", "volume", "page", "records_type",
            "nickname", "grantor", "interest", "grantee",
            "footnote_ref", "footnote_text", "full_text"]
HOT_WIDTHS = [8, 22, 8, 8, 26, 13, 13, 8, 8, 16, 20, 26, 30, 26, 9, 45, 60]

UNCLEAR_MARKERS = ("Unclear, see instrument", "Heirs of", "Estate of")


def write_history_of_title_xlsx(rows, out_path):
    font = "Arial"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name=font, size=10, bold=True, color="FFFFFF")
    body_font = Font(name=font, size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "History of Title"
    ws.append(HOT_HEADERS)
    for c in range(1, len(HOT_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 32

    date_cols = {HOT_KEYS.index("date_executed") + 1, HOT_KEYS.index("date_effective") + 1}
    for r in rows:
        ws.append([r.get(k, "") for k in HOT_KEYS])

    date_type = type(datetime.now().date())
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(HOT_HEADERS)):
        for cell in row_cells:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = border
        for dc in date_cols:
            cell = row_cells[dc - 1]
            if isinstance(cell.value, date_type):
                cell.number_format = "m/d/yyyy"
            cell.alignment = Alignment(horizontal="center", vertical="top")
        for key in ("chain", "item_no", "volume", "page", "footnote_ref"):
            row_cells[HOT_KEYS.index(key)].alignment = Alignment(horizontal="center", vertical="top")

    for i, w in enumerate(HOT_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)
    print(f"File generated: {out_path} ({len(rows)} records)")


def write_missing_fields_report(rows, out_path):
    """A plain-text audit of every row this script could not fill in with confidence --
    a blank/'Unclear' Grantor or Grantee, or a missing Book/Page -- so a reviewer can go
    straight to the source instrument instead of hunting through the full spreadsheet."""
    lines = []
    flagged = 0
    for r in rows:
        issues = []
        grantor = (r.get('grantor') or '').strip()
        grantee = (r.get('grantee') or '').strip()
        if not grantor or grantor.startswith(UNCLEAR_MARKERS):
            issues.append(f"Grantor: {grantor or '(empty)'}")
        if not grantee or grantee.startswith(UNCLEAR_MARKERS):
            issues.append(f"Grantee: {grantee or '(empty)'}")
        if not str(r.get('volume', '')).strip():
            issues.append("Book/Volume: (empty)")
        if not str(r.get('page', '')).strip():
            issues.append("Page: (empty)")
        if not r.get('instr_type'):
            issues.append("Instrument Type: not recognized")

        if issues:
            flagged += 1
            lines.append(f"Tract {r['tract_label']} - Item #{r['item_no']} (chain {r['chain']})")
            for issue in issues:
                lines.append(f"  - {issue}")
            lines.append(f"  Full Item Text: {r['full_text']}")
            lines.append("")

    header = [
        "MISSING / UNCERTAIN FIELDS REPORT",
        f"Total items: {len(rows)}   |   Items flagged: {flagged}",
        "=" * 70,
        "",
    ]
    content = "\n".join(header + lines) if lines else "\n".join(header + ["(No item was left incomplete.)"])
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Report generated: {out_path} ({flagged} items flagged)")


def main():
    if len(sys.argv) < 3:
        script = os.path.basename(sys.argv[0]) or "history_of_title_to_xlsx.py"
        print(f"Usage: python3 {script} input.(txt|pdf) output_prefix [county] [state]")
        print('  Accepts EITHER a plain-text export OR the PDF itself (auto-detected by extension).')
        print('  Generates <output_prefix>_HistoryOfTitle.xlsx, _CD.xlsx, _NCD.xlsx, and')
        print('  _MissingFields.txt (every row with a blank/uncertain field, for review).')
        print('  county/state are auto-detected from the document text; pass them explicitly')
        print('  only to override the detection (or to supply a value when nothing is found).')
        sys.exit(1)
    in_path, out_prefix = sys.argv[1], sys.argv[2]
    override_county = sys.argv[3] if len(sys.argv) > 3 else None
    override_state = sys.argv[4] if len(sys.argv) > 4 else None

    if in_path.lower().endswith('.pdf'):
        raw_text = pdf_to_indented_text(in_path)
        # save the extracted text alongside the input, purely so it can be inspected --
        # parsing always runs on the text held in memory, not this file
        txt_sidecar = f"{out_prefix}_extracted.txt"
        with open(txt_sidecar, "w", encoding="utf-8") as f:
            f.write(raw_text)
        print(f"Text extracted from the PDF saved to: {txt_sidecar}")
    else:
        with open(in_path, encoding="utf-8") as f:
            raw_text = f.read()

    detected_county, detected_state, how = detect_county_state(raw_text)
    county = override_county or detected_county or ""
    state = override_state or detected_state or ""

    if override_county or override_state:
        print(f"County/State: using command-line override -> {county}, {state} "
              f"(auto-detection found: {detected_county}, {detected_state})")
    elif detected_county:
        print(f"County/State: {county}, {state} ({how})")
    else:
        print(f"County/State: NOT detected in the document text ({how}). "
              f"Pass them explicitly: ... output_prefix \"County Name\" ST")

    tracts, footnotes = parse_document(raw_text)
    rows = build_records(tracts, footnotes)
    normalized = build_normalized_rows(rows, county, state)

    cd_records = [r for r in normalized if r["Class"] == "CD"]
    ncd_records = [r for r in normalized if r["Class"] == "NCD"]

    hot_path = f"{out_prefix}_HistoryOfTitle.xlsx"
    cd_path = f"{out_prefix}_CD.xlsx"
    ncd_path = f"{out_prefix}_NCD.xlsx"
    report_path = f"{out_prefix}_MissingFields.txt"

    write_history_of_title_xlsx(rows, hot_path)
    save_normalized_excel(cd_records, cd_path)
    save_normalized_excel(ncd_records, ncd_path)
    write_missing_fields_report(rows, report_path)

    total = len(rows)
    matched = sum(1 for r in rows if r['instr_type'])
    uncertain = sum(1 for r in rows if r['grantor'].startswith(UNCLEAR_MARKERS)
                     or r['grantee'].startswith(UNCLEAR_MARKERS))
    print(f"Tracts found: {len(tracts)}")
    for t in tracts:
        chains = sorted(set(it['chain'] for it in t['items']))
        print(f"  {t['label']}. {t['name']}: {len(t['items'])} items, chains={chains}")
    print(f"Footnotes found: {len(footnotes)}")
    print(f"Total items: {total}")
    print(f"Instrument type recognized: {matched}/{total}")
    print(f"Grantor/Grantee uncertain (Unclear/Estate of/Heirs of): {uncertain}/{total}")
    print(f"History of Title -> {hot_path}")
    print(f"CD records: {len(cd_records)} -> {cd_path}")
    print(f"NCD records: {len(ncd_records)} -> {ncd_path}")
    print(f"Missing-fields report -> {report_path}")


if __name__ == "__main__":
    main()
